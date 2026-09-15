import pandas as pd
from typing import Union
from openpyxl.utils import get_column_letter


class ExcelExporter:
    """Excel 내보내기 클래스"""
    
    def export(self, data: Union[pd.DataFrame, dict], output_path: str) -> None:
        """데이터를 Excel 파일로 저장"""
        if isinstance(data, dict):
            df = pd.DataFrame(data)
        elif isinstance(data, pd.DataFrame):
            df = data
        else:
            raise ValueError("지원하지 않는 데이터 형식입니다.")
        
        # Excel 파일로 저장
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Sheet1')
            
            # 워크시트 너비 자동 조정
            worksheet = writer.sheets['Sheet1']
            for idx, col in enumerate(df.columns):
                max_length = max(
                    df[col].astype(str).map(len).max(),
                    len(str(col))
                )
                # get_column_letter를 사용하여 26개 초과 컬럼 처리 (AA, AB 등)
                column_letter = get_column_letter(idx + 1)
                worksheet.column_dimensions[column_letter].width = min(max_length + 2, 50)
